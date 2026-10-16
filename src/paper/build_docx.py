r"""
Build the two IEIE 2026 fall submissions (v5.1) from the official templates:
  docs/paper/ieie2026f_proceedings_final.docx  (proceedings, double column, <= 5 pages, with authors)
  docs/paper/ieie2026f_blind_final.docx        (written review, journal template, 4 pages, no identifiers)
Text: src/paper/paper_text_v6.py (structure: method / dry results / discussion). Figures: docs/paper/figures/*.png (src/paper/make_figures.py).
Typography: number-unit spaces become U+00A0 (no stretch, no line break); figure/table/equation spacing set here.

Usage (repo root): python src\paper\build_docx.py [--suffix .new]
"""
from __future__ import annotations

import copy
import re
import shutil
import sys
import zipfile
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import paper_text_v6 as T  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
TPL = ROOT / "docs" / "paper" / "templates"
FIG = ROOT / "docs" / "paper" / "figures"
SUFFIX = ""
for i, a in enumerate(sys.argv):
    if a == "--suffix" and i + 1 < len(sys.argv):
        SUFFIX = sys.argv[i + 1]
OUT_P = ROOT / "docs" / "paper" / f"ieie2026f_proceedings_final{SUFFIX}.docx"
OUT_B = ROOT / "docs" / "paper" / f"ieie2026f_blind_final{SUFFIX}.docx"
COLW = Cm(7.95)
FIGW = {"fig1_architecture": 7.2, "fig2_transition": 7.0, "fig3_diffusion": 6.8, "fig4_window300": 7.6,
        "fig5_recess_spec": 7.2, "fig6_factors": 7.6}          # cm; <= column width
FIGW_B = {"fig1_architecture": 6.8, "fig2_transition": 7.6, "fig3_diffusion": 7.4, "fig4_window300": 7.6,
          "fig5_recess_spec": 7.6, "fig6_factors": 7.6}
GAP = dict(fig_before=10, cap_gap=4, cap_after=10, tab_before=10, tab_gap=4, tab_after=8, eq=8)

# proceedings only (the blind build never uses these)
AUTHOR_KO, AFFIL_KO, AUTHOR_EMAIL = "박주상", "성균관대학교", "pwqppt@gmail.com"
AUTHOR_EN, AFFIL_EN = "Joosang Park", "Department of AI Convergence, Sungkyunkwan University"

# ------------------------------------------------------------------ typography
NB = "\u00a0"
_UNITS = r"(?:°C|nm|µm|min|h|s|kcal/mol|GPa|ppm/K|ppm|%|K)"


def typo(s: str) -> str:
    """non-breaking, non-stretching spaces between numbers and units / symbols; Fig./Table numbers"""
    s = re.sub(r"(\d)[ ](" + _UNITS + r")(?![A-Za-z])", r"\1" + NB + r"\2", s)
    s = re.sub(r"(\d)[ ]×[ ](?=\d)", r"\1" + NB + "×" + NB, s)
    s = re.sub(r"([≤≥≈±<>])[ ](?=[\d−\-])", r"\1" + NB, s)
    s = re.sub(r"(\d)[ ]±", r"\1" + NB + "±", s)
    s = re.sub(r"\b(Fig\.|Table|Eq\.)[ ](?=\d)", r"\1" + NB, s)
    s = re.sub(r"(\d)[ ](?=[µn]m/)", r"\1" + NB, s)
    return s


# ------------------------------------------------------------------ citations
_CITE = re.compile(r"\[@([A-Z0-9]+(?:,@[A-Z0-9]+)*)\]")


class Cites:
    def __init__(self):
        self.order = []

    def scan(self, text):
        for m in _CITE.finditer(text):
            for k in m.group(1).split(",@"):
                if k not in self.order:
                    self.order.append(k)

    def num(self, text):
        def rep(m):
            ns = sorted(self.order.index(k) + 1 for k in m.group(1).split(",@"))
            return "[" + ",".join(str(n) for n in ns) + "]"
        return _CITE.sub(rep, text)

    def refs(self):
        return [(i + 1, T.REFS[k]) for i, k in enumerate(self.order)]


# ------------------------------------------------------------------ helpers
_MARK = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\}|\[\d{1,2}(?:[,–-]\d{1,2})*\]|<i>.*?</i>)")


def _runs(text):
    """yield (text, kind) where kind in {'', 'sub', 'sup', 'cite', 'i'}"""
    pos = 0
    for m in _MARK.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], ""
        t = m.group(0)
        if t.startswith("_{"):
            yield t[2:-1], "sub"
        elif t.startswith("^{"):
            yield t[2:-1], "sup"
        elif t.startswith("<i>"):
            yield t[3:-4], "i"
        else:
            yield t, "cite"
        pos = m.end()
    if pos < len(text):
        yield text[pos:], ""


def _font(run, name, size=None, bold=None, east=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi"):
        rf.set(qn(a), name)
    rf.set(qn("w:eastAsia"), east or name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def fill(par, text, font, size=None, east=None, bold=None, cite=True):
    for t, kind in _runs(text):
        if kind == "cite" and not cite:
            kind = ""
        r = par.add_run(t)
        _font(r, font, size, bold, east)
        if kind in ("sub",):
            r.font.subscript = True
        if kind in ("sup", "cite"):
            r.font.superscript = True
        if kind == "i":
            r.font.italic = True
    return par


def clear_after(body, keep_upto_el):
    """remove every body child after keep_upto_el except the final sectPr"""
    seen = False
    for el in list(body.iterchildren()):
        if el is keep_upto_el:
            seen = True; continue
        if seen and el.tag != qn("w:sectPr"):
            body.remove(el)


def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side, val in (("top", top), ("bottom", bottom), ("left", None), ("right", None)):
        e = OxmlElement(f"w:{side}")
        if val:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:color"), "000000")
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)


def add_table(doc, spec, font="Times New Roman", size=7.0, before=None):
    cap, head, rows, widths = spec
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    mar = OxmlElement("w:tblCellMar")
    for side in ("left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "45"); e.set(qn("w:type"), "dxa"); mar.append(e)
    tblPr.append(mar)
    for ri, row in enumerate([head] + rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci); c.width = Cm(widths[ci])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            fill(p, typo(val), font, size, bold=(ri == 0), cite=True)
            set_cell_borders(c, top=8 if ri == 0 else (4 if ri == 1 else None),
                             bottom=8 if ri == len(rows) else None)
    return t



def _single_column_break(doc, after_el):
    """insert a paragraph carrying a one-column sectPr after after_el; make the final section continuous"""
    final = doc.element.body.find(qn("w:sectPr"))
    sp = copy.deepcopy(final)
    cols = sp.find(qn("w:cols"))
    if cols is not None:
        for a in list(cols.attrib):
            del cols.attrib[a]
        cols.set(qn("w:space"), "425")
    p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr"); ppr.append(sp); p.append(ppr)
    after_el.addnext(p)
    t = final.find(qn("w:type"))
    if t is None:
        t = OxmlElement("w:type"); final.insert(0, t)
    t.set(qn("w:val"), "continuous")
    return p




ROMAN_P = ["I", "II", "III", "IV", "V"]
ROMAN_B = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ"]


def _space(p, before=None, after=None):
    pf = p.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)


def resolve(text, ver, fign, tabn, eqn=None):
    """version tags, figure/table/equation references"""
    eqn = eqn or {}
    text = re.sub(r"\{\{P:(.*?)\}\}", lambda m: m.group(1) if ver == "P" else "", text, flags=re.S)
    text = re.sub(r"\{\{B:(.*?)\}\}", lambda m: m.group(1) if ver == "B" else "", text, flags=re.S)
    text = re.sub(r"\[\[FIG:(\w+)\]\]", lambda m: f"(Fig. {fign[m.group(1)]})" if m.group(1) in fign else "", text)
    text = re.sub(r"\[\[FIGN:(\w+)\]\]", lambda m: f"Fig. {fign[m.group(1)]}", text)
    text = re.sub(r"\[\[TAB:(\w+)\]\]", lambda m: f"Table {tabn[m.group(1)]}", text)
    text = re.sub(r"\[\[EQ:(\w+)\]\]", lambda m: f"({eqn[m.group(1)]})", text)
    return text


def version_items(ver):
    """flatten T.BODY for one version: [(kind, payload)] with resolved text, numbered headings, equations and floats"""
    figs = T.FIGS_P if ver == "P" else T.FIGS_B
    tabs = T.TABS_P if ver == "P" else T.TABS_B
    fign = {k: i + 1 for i, k in enumerate(figs)}
    tabn = {k: i + 1 for i, k in enumerate(tabs)}
    body = []
    for kind, x in T.BODY:
        if kind in ("PP", "PB", "H2P", "H2B", "PCP", "PCB", "EQP", "EQB", "FIGP", "FIGB"):
            if kind[-1] != ver:
                continue
            kind = kind[:-1]
        body.append((kind, x))
    eqn = {}
    for kind, x in body:
        if kind == "EQ" and x[0] not in eqn:
            eqn[x[0]] = len(eqn) + 1
    out, h1, h2 = [], 0, 0
    for kind, x in body:
        if kind == "H1":
            h1 += 1; h2 = 0
            if ver == "P":
                out.append(("h1", f"{ROMAN_P[h1 - 1]}. {x}"))
            else:
                t = "  ".join(x) if len(x) == 2 else x
                out.append(("h1", f"{ROMAN_B[h1 - 1]}. {t}"))
        elif kind == "H2":
            h2 += 1
            out.append(("h2", f"{h1}.{h2} {x}" if ver == "P" else f"{h2}. {x}"))
        elif kind in ("P", "PC"):
            out.append((kind.lower(), resolve(x, ver, fign, tabn, eqn)))
        elif kind == "EQ":
            out.append(("eq", (x[1], eqn[x[0]])))
        elif kind == "FIG":
            if x in fign:
                out.append(("fig", (T.FIGFILE[x], f"Fig. {fign[x]}. " + resolve(T.FIGCAP[x], ver, fign, tabn, eqn))))
        elif kind == "TAB":
            if x in tabn:
                cap, head, rows, w = T.TABLES[x]
                rows = [[resolve(c, ver, fign, tabn, eqn) for c in r] for r in rows]
                out.append(("tab", (f"Table {tabn[x]}. " + cap, head, rows, w)))
    return out


def _eq_tabs(p):
    """centred equation with right-aligned number"""
    ts = p.paragraph_format.tab_stops
    ts.add_tab_stop(Cm(COLW.cm / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(COLW - Cm(0.05), WD_TAB_ALIGNMENT.RIGHT)


def _texts(item):
    kind, x = item
    if kind in ("p", "pc"):
        return [x]
    if kind == "tab":
        return [x[0]] + [c for r in x[2] for c in r]
    if kind == "fig":
        return [x[1]]
    return []


def _cites(items):
    ct = Cites()
    for it in items:
        for t in _texts(it):
            ct.scan(t)
    return ct


# ------------------------------------------------------------------ proceedings
def build_proceedings():
    items = version_items("P")
    ct = _cites(items)
    d = docx.Document(str(TPL / "논문양식.docx"))
    xbody = d.element.body
    t_title, t_auth = d.tables[0], d.tables[1]
    ps = [p for c in t_title._cells for p in c.paragraphs if p.text.strip()]
    groups = [T.TITLE_KO_LINES[:1], T.TITLE_KO_LINES[1:]]
    for p, lines in zip(ps, groups):
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
        r0 = p.runs[0]; r0.text = lines[0]; r0.font.size = Pt(18)
        for extra in lines[1:]:
            r0.add_break(); r0.add_text(extra)
    for c in t_auth._cells:
        for p in c.paragraphs:
            tx = p.text
            if not tx.strip():
                continue
            new = None
            if tx.strip().startswith("***,"):
                new = AUTHOR_KO
            elif "소속" in tx:
                new = AFFIL_KO
            elif tx.strip().startswith("e-mail"):
                new = "e-mail : " + AUTHOR_EMAIL
            elif tx.startswith("Implementation"):
                new = T.TITLE_EN_LINES[0]
            elif tx.startswith("Using"):
                new = T.TITLE_EN_LINES[1]
            elif "and" in tx and "***" in tx:
                new = AUTHOR_EN
            elif "University" in tx:
                new = AFFIL_EN
            if new is not None:
                for r in p.runs[1:]:
                    r._element.getparent().remove(r._element)
                p.runs[0].text = new
                if new in T.TITLE_EN_LINES:
                    p.runs[0].font.size = Pt(11)
    sect_p = next(el for el in xbody.iterchildren() if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None)
    clear_after(xbody, sect_p)
    F, FS = "HY신명조", 9

    def para(text, size=FS, align="both", indent=True, before=0, after=0, ls=1.2):
        p = d.add_paragraph(style="바탕글")
        pf = p.paragraph_format; pf.line_spacing = ls; pf.space_after = Pt(after); pf.space_before = Pt(before)
        p.alignment = {"both": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        if indent:
            pf.first_line_indent = Pt(9)
        fill(p, typo(ct.num(text)), "Times New Roman", size, east=F)
        return p

    para("Abstract", size=12, align="center", indent=False).runs[0].font.bold = True
    para(T.ABSTRACT_EN, after=2)
    for kind, x in items:
        if kind == "h1":
            para(x, size=12, align="center", indent=False, before=6, after=2).paragraph_format.keep_with_next = True
        elif kind == "h2":
            p = para(x, align="left", indent=False, before=6, after=1); p.paragraph_format.keep_with_next = True
        elif kind == "p":
            para(x)
        elif kind == "pc":
            para(x, indent=False)
        elif kind == "eq":
            line, n = x
            eq = para("", indent=False, align="left", before=GAP["eq"], after=GAP["eq"], ls=1.0)
            _eq_tabs(eq)
            fill(eq, "\t" + typo(line).replace(" ", NB) + f"\t({n})", "Times New Roman", 9, east=F)
        elif kind == "fig":
            name, cap = x
            p = d.add_paragraph(style="바탕글"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format; pf.space_before = Pt(GAP["fig_before"]); pf.space_after = Pt(0); pf.line_spacing = 1.0
            pf.keep_with_next = True
            p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(FIGW[name]))
            para(cap, size=9, align="center", indent=False, before=GAP["cap_gap"], after=GAP["cap_after"], ls=1.0)
        elif kind == "tab":
            c = para(x[0], size=9, align="center", indent=False, before=GAP["tab_before"], after=GAP["tab_gap"], ls=1.0)
            c.paragraph_format.keep_with_next = True
            spec = (x[0], x[1], [[ct.num(v) for v in r] for r in x[2]], x[3])
            t = add_table(d, spec); _no_split(t)
            para("", size=2, indent=False, ls=1.0, after=GAP["tab_after"])
    para("참고문헌", size=12, align="center", indent=False, before=6, after=2)
    for n, r in ct.refs():
        p = d.add_paragraph(style="바탕글")
        pf = p.paragraph_format; pf.line_spacing = 1.2; pf.left_indent = Pt(15.4); pf.first_line_indent = Pt(-15.4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fill(p, f"[{n}] ", "Times New Roman", 9, east=F, cite=False)
        fill(p, typo(r), "Times New Roman", 9, east=F, cite=False)
    cp = d.core_properties
    for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
        setattr(cp, k, "")
    _strip_caption_periods(d)
    d.save(str(OUT_P))
    _scrub_app(OUT_P)
    return ct


# ------------------------------------------------------------------ blind (journal template)
def build_blind():
    items = version_items("B")
    ct = _cites(items)
    d = docx.Document(str(TPL / "서면 심사용 논문 양식.docx"))
    xbody = d.element.body
    hdr = d.tables[0]

    def set_text(p, text, cite=False):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        fill(p, typo(text), "Times New Roman", None, east="한양신명조", cite=cite)
    seen_body, done, flags = 0, set(), set()
    for c in hdr._cells:
        if id(c._tc) in done:
            continue
        done.add(id(c._tc))
        for p in list(c.paragraphs):
            n = p.style.name; tx = p.text.strip()
            if n == "논문번호":
                if "label" not in flags:
                    set_text(p, "학부생 논문경진대회 서면 심사용"); flags.add("label")
                else:
                    set_text(p, "")
            elif n == "한글제목":
                set_text(p, T.TITLE_KO)
            elif n == "영문제목":
                set_text(p, "(" + T.TITLE_EN + ")")
            elif n == "요약본문":
                seen_body += 1
                set_text(p, T.ABSTRACT_KO if seen_body == 1 else T.ABSTRACT_EN_B)
            elif "Keywords" in tx:
                set_text(p, "      Keywords : " + T.KEYWORDS)
    tblPr = hdr._tbl.tblPr
    for e in tblPr.findall(qn("w:tblpPr")):
        tblPr.remove(e)
    for trPr in hdr._tbl.iter(qn("w:trPr")):
        for e in trPr.findall(qn("w:tblHeader")):
            trPr.remove(e)
    for tr in list(hdr._tbl.findall(qn("w:tr")))[1:]:
        if not "".join(t.text or "" for t in tr.iter(qn("w:t"))).strip():
            hdr._tbl.remove(tr)
    clear_after(xbody, hdr._tbl)
    brk = _single_column_break(d, hdr._tbl)
    ppr = brk.find(qn("w:pPr")); sp_ = OxmlElement("w:spacing"); sp_.set(qn("w:before"), "0"); sp_.set(qn("w:after"), "240")
    sp_.set(qn("w:line"), "240"); sp_.set(qn("w:lineRule"), "auto"); ppr.insert(0, sp_)

    def p_(style, text, cite=True):
        p = d.add_paragraph(style=style)
        fill(p, typo(ct.num(text)), "Times New Roman", None, east="한양신명조", cite=cite)
        if style == "Body Text":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return p
    for kind, x in items:
        if kind == "h1":
            p_("각 장 제목", x).paragraph_format.keep_with_next = True
        elif kind == "h2":
            p_("각 절 제목", x).paragraph_format.keep_with_next = True
        elif kind == "p":
            p_("Body Text", x)
        elif kind == "pc":
            p = p_("Body Text", x); p.paragraph_format.first_line_indent = Pt(0)
        elif kind == "eq":
            line, n = x
            eq = d.add_paragraph(style="수식번호")
            pf = eq.paragraph_format; pf.first_line_indent = Pt(0); pf.left_indent = Pt(0); pf.line_spacing = 1.0
            eq.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _eq_tabs(eq); _space(eq, GAP["eq"], GAP["eq"])
            fill(eq, "\t" + typo(line).replace(" ", NB) + f"\t({n})", "Times New Roman", 10, east="한양신명조")
        elif kind == "fig":
            name, cap = x
            p = d.add_paragraph(style="Body Text"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format; pf.first_line_indent = Pt(0); pf.line_spacing = 1.0; pf.keep_with_next = True
            _space(p, GAP["fig_before"], 0)
            p.add_run().add_picture(str(FIG / f"{name}.png"), width=Cm(FIGW_B[name]))
            c = p_("그림/표 제목", cap); _space(c, GAP["cap_gap"], GAP["cap_after"])
        elif kind == "tab":
            c = p_("그림/표 제목", x[0]); c.paragraph_format.keep_with_next = True; _space(c, 6, GAP["tab_gap"])
            spec = (x[0], x[1], [[ct.num(v) for v in r] for r in x[2]], x[3])
            t = add_table(d, spec, size=7.0); _no_split(t)
            sp = d.add_paragraph(style="Body Text"); sp.paragraph_format.line_spacing = 0.6; _space(sp, 0, GAP["tab_after"])
    p_("각 장 제목", "REFERENCES")
    for n, r in ct.refs():
        p_("참고문헌내용", f"[{n}] " + r, cite=False)
    cp = d.core_properties
    for k in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "category"):
        setattr(cp, k, "")
    _strip_caption_periods(d)
    d.save(str(OUT_B))
    _scrub_app(OUT_B)
    return ct


def _no_split(t):
    """keep table rows together (cantSplit + keep_with_next on all but last row)"""
    rows = t.rows
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit"); trPr.append(e)
        if i < len(rows) - 1:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True


def _scrub_app(path):
    """blank Company/Manager in docProps/app.xml (identifying metadata)"""
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                s = data.decode("utf-8")
                s = re.sub(r"<Company>.*?</Company>", "<Company></Company>", s, flags=re.S)
                s = re.sub(r"<Manager>.*?</Manager>", "<Manager></Manager>", s, flags=re.S)
                data = s.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(str(tmp), str(path))


_CAP_RE = re.compile(r"^(Fig\.|Table)\s?\d+\.\s")


def _strip_caption_periods(d):
    """Template rule '(마침표삭제)': no full stop at the end of a figure/table caption."""
    n = 0
    for p in d.paragraphs:
        if _CAP_RE.match(p.text) and p.text.rstrip().endswith("."):
            for r in reversed(p.runs):
                if r.text.strip():
                    r.text = r.text.rstrip().rstrip(".")
                    n += 1
                    break
    return n


if __name__ == "__main__":
    build_proceedings(); print("ok", OUT_P.name)
    build_blind(); print("ok", OUT_B.name)
